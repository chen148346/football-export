"""
fixed_exporter.py - 固定表头 Excel 导出模块
==============================================
将完场比赛导出为 **表头固定** 的 8-Sheet Excel 文件。

8 个 Sheet（固定顺序、固定中文表头）：
1. match_overview     比赛
2. tech_stats         指标
3. near_matches       近绩
4. vs_matches         往绩
5. first_half_stats   上半场统计
6. second_half_stats  下半场统计
7. league_rank_stats  联赛排名
8. detail_events      详细事件

核心设计：
- 表头固定：与固定 schema（fixed_schema.py）一致，不随比赛动态增列。
- 双编码归一：指标/上下半场/联赛排名中的统计项，英文键与数字码归一到同一中文列。
- 简体中文：全部字段统一简体中文。
- 球队名标准化：通过 TeamMaster 按 team_id 归一，未知 ID 回退原始名。
- 缺失字段留空。
"""

import os
import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from . import config
from . import json_parser
from . import fixed_schema
from .team_master import get_team_master

# ============================================================================
# 样式（与 V1.7 保持一致）
# ============================================================================
HEADER_FONT = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATA_FONT = Font(name="Microsoft YaHei", size=10)
DATA_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=False)
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


def _beijing_now() -> datetime.datetime:
    utc_now = datetime.datetime.utcnow()
    return utc_now + datetime.timedelta(hours=config.BEIJING_TZ_OFFSET_HOURS)


def _fmt(value):
    """格式化单元格值：None 与未命中统一留空。"""
    if value is None:
        return ""
    return value


class FixedExporter:
    """固定表头导出器：持有 TeamMaster 并构建 8 个固定 Sheet。"""

    def __init__(self, team_master=None):
        self.tm = team_master or get_team_master()

    # ---------------- 比赛 Sheet ----------------
    def build_match_sheet(self, matches) -> (list, list):
        headers = fixed_schema.MATCH_FIXED_HEADERS
        rows = []
        for m in matches:
            home_formation = away_formation = ""
            home_corner = away_corner = home_half_corner = away_half_corner = ""
            referee_name = ""
            if m.shijian_json:
                lineup = json_parser.extract_lineup(m.shijian_json)
                home_formation = lineup["home_formation"]
                away_formation = lineup["away_formation"]
                corner = json_parser.extract_corner_events(m.shijian_json)
                home_corner = corner["home_corner"]
                away_corner = corner["away_corner"]
                home_half_corner = corner["home_half_corner"]
                away_half_corner = corner["away_half_corner"]
            if m.analysis_json:
                ref = json_parser.extract_referee(m.analysis_json)
                referee_name = ref["referee_name"]

            match_time_dt = json_parser.parse_match_time_14(m.match_time)
            match_time_str = json_parser.format_datetime(match_time_dt)

            home_team = self.tm.canonical_by_variant(m.home_team) or m.home_team
            away_team = self.tm.canonical_by_variant(m.away_team) or m.away_team

            rows.append({
                "match_id": m.match_id,
                "sclass_name": m.sclass_name,
                "match_date": m.match_date,
                "match_time": match_time_str,
                "home_team": home_team,
                "away_team": away_team,
                "home_rank": m.home_rank,
                "away_rank": m.away_rank,
                "home_score": m.home_score,
                "away_score": m.away_score,
                "home_half_score": m.home_half_score,
                "away_half_score": m.away_half_score,
                "home_red": m.home_red,
                "away_red": m.away_red,
                "home_yellow": m.home_yellow,
                "away_yellow": m.away_yellow,
                "home_formation": home_formation,
                "away_formation": away_formation,
                "home_corner": home_corner,
                "away_corner": away_corner,
                "home_half_corner": home_half_corner,
                "away_half_corner": away_half_corner,
                "referee_name": referee_name,
            })
        return headers, rows

    # ---------------- 指标 / 上下半场统计 ----------------
    def _build_stat_sheet(self, matches, list_key, fixed_cn_list):
        """通用：根据固定统计项中文清单，构建指标/上下半场 sheet。"""
        zhu_header = [f"主队_{cn}" for cn in fixed_cn_list]
        ke_header = [f"客队_{cn}" for cn in fixed_cn_list]
        headers = ["比赛编号"] + zhu_header + ke_header

        rows = []
        for m in matches:
            # 统计项容器：中文名 -> {"home": val, "away": val}
            stat_by_cn = {}
            if m.shijian_json:
                tech = json_parser.extract_tech_stats_flat(
                    m.shijian_json, m.match_id, list_key=list_key
                )
                stat_by_cn = tech["stats_cn"]
            row = {"比赛编号": m.match_id}
            for cn in fixed_cn_list:
                info = stat_by_cn.get(cn, {})
                row[f"主队_{cn}"] = _fmt(info.get("home"))
                row[f"客队_{cn}"] = _fmt(info.get("away"))
            rows.append(row)
        return headers, rows

    def build_tech_sheet(self, matches):
        return self._build_stat_sheet(matches, "itemList", fixed_schema.TECH_FIXED_CN)

    def build_first_half_sheet(self, matches):
        return self._build_stat_sheet(matches, "firstHalfList",
                                      fixed_schema.FIRST_HALF_FIXED_CN)

    def build_second_half_sheet(self, matches):
        return self._build_stat_sheet(matches, "secondHalfList",
                                      fixed_schema.SECOND_HALF_FIXED_CN)

    # ---------------- 近绩 / 往绩 ----------------
    def _build_brief_sheet(self, matches, mode: str):
        headers = fixed_schema.BRIEF_FIXED_HEADERS
        rows = []
        for m in matches:
            if mode == "near":
                briefs = json_parser.extract_near_matches(m.analysis_json, m.match_id) \
                    if m.analysis_json else []
            else:  # vs
                briefs = json_parser.extract_vs_matches(m.analysis_json, m.match_id) \
                    if m.analysis_json else []
            for b in briefs:
                home_name = self.tm.canonical_name(b["home_team_id"], b["home_team_name"])
                away_name = self.tm.canonical_name(b["away_team_id"], b["away_team_name"])
                h, a = b["home_score"], b["away_score"]
                result = self._result_text(h, a)
                rows.append({
                    "match_id": b["match_id"],
                    "match_time": b["match_time"],
                    "league_name": b["league_name"],
                    "home_team": home_name,
                    "home_id": b["home_team_id"],
                    "home_score": h,
                    "home_half_score": b["home_half_score"],
                    "home_corner": b["home_corner"],
                    "home_yellow": b["home_yellow"],
                    "home_red": b["home_red"],
                    "away_team": away_name,
                    "away_id": b["away_team_id"],
                    "away_score": a,
                    "away_half_score": b["away_half_score"],
                    "away_corner": b["away_corner"],
                    "away_yellow": b["away_yellow"],
                    "away_red": b["away_red"],
                    "result": result,
                })
                # 增量学习球队名称
                self.tm.observe(b["home_team_id"], b["home_team_name"])
                self.tm.observe(b["away_team_id"], b["away_team_name"])
        return headers, rows

    def build_near_sheet(self, matches):
        return self._build_brief_sheet(matches, "near")

    def build_vs_sheet(self, matches):
        return self._build_brief_sheet(matches, "vs")

    # ---------------- 联赛排名 ----------------
    def build_league_rank_sheet(self, matches):
        headers = (["比赛编号", "主队"]
                   + [f"主队_{cn}" for cn in fixed_schema.LEAGUE_RANK_FIXED]
                   + ["客队"] + [f"客队_{cn}" for cn in fixed_schema.LEAGUE_RANK_FIXED])
        rows = []
        for m in matches:
            rank = {}
            if m.analysis_json:
                rank = json_parser.extract_league_rank_stats_cn(m.analysis_json, m.match_id)
            home_name = self.tm.canonical_by_variant(m.home_team) or m.home_team
            away_name = self.tm.canonical_by_variant(m.away_team) or m.away_team
            row = {"比赛编号": m.match_id, "主队": home_name, "客队": away_name}
            for cn, kind in fixed_schema.LEAGUE_RANK_FIXED.items():
                info = rank.get(kind, {})
                row[f"主队_{cn}"] = _fmt(info.get("home"))
                row[f"客队_{cn}"] = _fmt(info.get("away"))
            rows.append(row)
        return headers, rows

    # ---------------- 详细事件 ----------------
    def build_detail_events_sheet(self, matches):
        headers = fixed_schema.FIXED_SHEETS[7][2]
        rows = []
        for m in matches:
            events = json_parser.extract_detailed_events_fixed(m.shijian_json, m.match_id) \
                if m.shijian_json else []
            for ev in events:
                rows.append({
                    "比赛编号": ev["match_id"],
                    "事件时间": ev["time"],
                    "事件类型": ev["event_type"],
                    "球队": ev["team"],
                    "球员": ev["player"],
                    "详情": ev["detail"],
                })
        return headers, rows

    # ---------------- 工具 ----------------
    @staticmethod
    def _result_text(h, a):
        try:
            h, a = int(h), int(a)
        except (TypeError, ValueError):
            return ""
        if h > a:
            return "主胜"
        if h < a:
            return "客胜"
        return "平"

    # ---------------- 总入口 ----------------
    def build_all(self, matches) -> dict:
        """返回 {sheet_id: (headers, rows)}，sheet 顺序按固定清单。"""
        builders = [
            ("match_overview", self.build_match_sheet),
            ("tech_stats", self.build_tech_sheet),
            ("near_matches", self.build_near_sheet),
            ("vs_matches", self.build_vs_sheet),
            ("first_half_stats", self.build_first_half_sheet),
            ("second_half_stats", self.build_second_half_sheet),
            ("league_rank_stats", self.build_league_rank_sheet),
            ("detail_events", self.build_detail_events_sheet),
        ]
        out = {}
        for sid, builder in builders:
            out[sid] = builder(matches)
        return out


# ============================================================================
# 对外导出函数
# ============================================================================

def export_fixed_excel(matches, output_path: str, teams=None) -> str:
    """
    将比赛导出为固定表头 8-Sheet Excel。

    参数：
        matches: MatchRecord 列表（需含 shijian_json/analysis_json）
        output_path: 输出文件路径
        teams: 可选自定义 TeamMaster 实例
    """
    exporter = FixedExporter(team_master=teams)
    sheets = exporter.build_all(matches)

    wb = Workbook()
    wb.remove(wb.active)
    title_map = {sid: title for sid, title, _ in fixed_schema.FIXED_SHEETS}
    for sid, (headers, rows) in sheets.items():
        ws = wb.create_sheet(title=title_map[sid])
        _write_fixed_sheet(ws, headers, rows)
    wb.save(output_path)
    return output_path


def _write_fixed_sheet(ws, headers, rows):
    """写固定表头 + 数据行（复用 V1.7 样式）。"""
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, header in enumerate(headers, 1):
            val = row_data.get(header)
            val = "" if val is None else val
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = DATA_FONT
            cell.alignment = DATA_ALIGN
            cell.border = THIN_BORDER
    _auto_column_width(ws, len(headers))
    ws.freeze_panes = "A2"


def _auto_column_width(ws, col_count: int, max_width: int = 40):
    for col in range(1, col_count + 1):
        max_len = 0
        col_letter = get_column_letter(col)
        for row in range(1, min(ws.max_row, 100) + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                length = sum(2 if ord(c) > 127 else 1 for c in str(val))
                if length > max_len:
                    max_len = length
        width = min(max_len + 2, max_width)
        ws.column_dimensions[col_letter].width = max(width, 8)


def generate_fixed_filename(league: str, start_date: str, end_date: str) -> str:
    """生成固定表头导出文件名（北京时间时间戳）。"""
    now = _beijing_now()
    ts = now.strftime("%Y%m%d_%H%M%S")
    league_part = league if league else "multi"
    return f"fixed_{league_part}_{start_date}_{end_date}_{ts}.xlsx"